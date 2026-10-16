"""Bounded, read-only CSV/XLSX parsing. Validate all rows before writing any SKU."""
import csv
import io
import re
import unicodedata
import zipfile
from decimal import Decimal, InvalidOperation

MAX_ROWS = 10000
HEADERS = ('sku', 'largura', 'altura', 'comprimento', 'peso bruto')


class ImportValidationError(ValueError):
    def __init__(self, errors):
        self.errors = errors[:50]
        super().__init__('Planilha inválida')


def header(value):
    value = unicodedata.normalize('NFKD', str(value or ''))
    value = ''.join(c for c in value if not unicodedata.combining(c)).lower().strip()
    return re.sub(r'\s*\((cm|kg)\)\s*$', '', value).replace('_', ' ')


def parse(upload, allowed_skus):
    content = upload.stream.read(2 * 1024 * 1024 + 1)
    if len(content) > 2 * 1024 * 1024:
        raise ImportValidationError(['O arquivo deve ter no máximo 2 MB.'])
    extension = (upload.filename or '').rsplit('.', 1)[-1].lower()
    try:
        if extension == 'csv':
            try:
                text = content.decode('utf-8-sig')
            except UnicodeDecodeError:
                text = content.decode('cp1252')
            try:
                dialect = csv.Sniffer().sniff(text[:8192], delimiters=';,\t')
            except csv.Error:
                dialect = csv.excel
                dialect.delimiter = ';'
            rows = csv.reader(io.StringIO(text), dialect)
            return validate(rows, allowed_skus)
        if extension == 'xlsx':
            from openpyxl import load_workbook
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries) > 2000 or sum(e.file_size for e in entries) > 20 * 1024 * 1024:
                    raise ImportValidationError(['Arquivo Excel muito grande após descompactar.'])
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
            try:
                if len(book.worksheets) != 1:
                    raise ImportValidationError(['Use um arquivo com apenas uma aba de dados.'])
                sheet = book.worksheets[0]
                if sheet.max_column and sheet.max_column > 20:
                    raise ImportValidationError(['Use somente as cinco colunas do modelo.'])
                return validate(sheet.iter_rows(values_only=True), allowed_skus)
            finally:
                book.close()
        raise ImportValidationError(['Envie um arquivo .xlsx ou .csv.'])
    except ImportValidationError:
        raise
    except Exception:
        raise ImportValidationError(['Não foi possível ler o arquivo. Confira o formato e salve novamente.']) from None


def validate(rows, allowed_skus):
    rows = iter(rows)
    names = next(rows, ())
    if tuple(header(n) for n in names) != HEADERS:
        raise ImportValidationError(['Cabeçalho obrigatório: sku;largura;altura;comprimento;peso bruto. Medidas em cm e peso em kg.'])
    entries, errors, seen = [], [], set()
    for line, row in enumerate(rows, 2):
        if line > MAX_ROWS + 1:
            raise ImportValidationError([f'Limite de {MAX_ROWS} linhas de dados por arquivo.'])
        if all(v is None or str(v).strip() == '' for v in row):
            continue
        if len(row) != 5:
            errors.append(f'Linha {line}: informe exatamente cinco colunas.')
            continue
        raw_sku = row[0]
        sku = str(raw_sku or '').strip()
        if sku not in allowed_skus and sku.startswith("'") and sku[1:] in allowed_skus:
            sku = sku[1:]
        if not isinstance(raw_sku, str):
            errors.append(f'Linha {line}: formate o SKU como texto para preservar zeros à esquerda.')
            continue
        if sku not in allowed_skus:
            errors.append(f'Linha {line}: SKU {sku[:80]} não localizado na conta. Confira o código ou atualize os anúncios.')
            continue
        if sku in seen:
            errors.append(f'Linha {line}: SKU {sku} duplicado no arquivo.')
            continue
        seen.add(sku)
        entry = {'sku': sku}
        try:
            for key, raw, places in zip(('width', 'height', 'length', 'weight'), row[1:], (2, 2, 2, 3)):
                value = Decimal(str(raw).strip().replace(',', '.'))
                if not value.is_finite() or value <= 0 or value > 10000 or value != value.quantize(Decimal(10) ** -places):
                    raise ValueError()
                entry[key] = format(value, 'f')
        except (ValueError, InvalidOperation):
            errors.append(f'Linha {line}: medidas e peso devem ser positivos, até 10.000; use até 2 casas nas medidas e 3 no peso. Fórmulas não são aceitas.')
            continue
        entries.append(entry)
    if not entries and not errors:
        errors.append('O arquivo não contém SKUs para importar.')
    if errors:
        raise ImportValidationError(errors)
    return entries
